"""
Stage 2 - human annotation (independent ground truth).

Answers R2.3, R3.2 (circularity / no ground truth), R4.3 (annotation strategy,
labelling criteria, validation), and enables R2.7 (error analysis).

build_sheets(): writes
  data/annotation_set.csv                    which reviews are in the gold set
  annotation/annotation_coder_A.xlsx         blind coding workbook, coder A order
  annotation/annotation_coder_B.xlsx         same items, different random order
  annotation/CODEBOOK.md                     coding instructions (also a sheet in each workbook)
Coders work independently and save their completed files into
annotation/returned/ with the same file names.

agreement(): reads the returned workbooks and writes
  results/04_human_annotation/inter_rater_agreement.csv
  annotation/adjudication.xlsx              disagreements to resolve together
  data/gold_labels.csv                      gold standard used in evaluation
"""
from __future__ import annotations

import json

import numpy as np
import pandas as pd

import config as C

OUT = C.RESULTS_DIR / "04_human_annotation"

EMOTIONS = ["REGRET", "DISAPPOINTMENT", "BOTH", "NEITHER"]
FORCED = ["REGRET", "DISAPPOINTMENT"]
FOCUS = ["BEHAVIOR", "SITUATION", "NONE"]
DIMS = ["counterfactual", "responsibility", "controllability", "affect"]
SCALE = [str(i) for i in range(1, 8)]

CODEBOOK = """# Coding manual: regret vs disappointment in customer reviews

You will read low-rated Amazon product reviews and code the emotion the reviewer
expresses. Work **independently**: do not discuss items with the other coder until
both of you have finished. Star ratings and all model outputs are deliberately hidden.

## Before you start (15 minutes, together)
Both coders label the 10 reviews in the `practice` sheet on their own, then compare and
discuss them using this manual. The practice reviews are NOT part of the study. After that,
code the `coding` sheet alone.

## Theoretical basis
Zeelenberg et al. (1998, 2000); Zeelenberg & Pieters (2007); Marcatto & Ferrante (2008).
Regret and disappointment are both negative, outcome-related emotions. They differ in
**who is responsible** for the bad outcome.

| | REGRET | DISAPPOINTMENT |
|---|---|---|
| Responsibility | The reviewer's own choice or action | The product, seller, brand or circumstances |
| Typical thought | "I should have bought X instead", "I wish I had read the reviews" | "It should have lasted longer", "It did not work as advertised" |
| Counterfactual | Changes own **behaviour** | Changes the **situation** |
| Control | Outcome was controllable by the reviewer | Outcome was outside the reviewer's control |

## Fields (REQUIRED fields are marked)
1. **primary_emotion** (REQUIRED): REGRET / DISAPPOINTMENT / BOTH / NEITHER.
   - BOTH: clear self-blame **and** clear blame of the product, with neither dominant.
   - NEITHER: no regret or disappointment, e.g. a neutral description, a mostly positive
     3-star review, a question, or spam.
2. **forced_choice** (REQUIRED): if you had to choose only REGRET or DISAPPOINTMENT, which one?
   Always fill this in, even when primary_emotion is BOTH or NEITHER.
3. **certainty** (REQUIRED): your confidence in the primary_emotion code (LOW / MEDIUM / HIGH).
4. **cf_focus** (recommended): primary focus of any counterfactual thinking: BEHAVIOR (own actions),
   SITUATION (product/circumstances), or NONE (no counterfactual expressed).
5. **counterfactual / responsibility / controllability / affect** (recommended, 1-7 each):
   - counterfactual: how much "what if / if only / should have" reasoning (1 = none, 7 = extensive)
   - responsibility: who is held responsible (1 = entirely external, 7 = entirely the reviewer)
   - controllability: how controllable the outcome was for the reviewer (1 = uncontrollable, 7 = fully)
   - affect: intensity of negative emotion (1 = mild, 7 = intense)
6. **notes** (optional): flag sarcasm, mixed reviews, non-English text, etc.

## Decision rules
- Code what the text **expresses**, not what you think the reviewer *should* feel.
- Product defects, wrong/damaged item, poor quality, short battery life, or "not as described"
  = DISAPPOINTMENT.
- **"Would not buy again" / "Won't buy this brand again" / "Stick with Duracell"** on their own are
  statements of future intention, not self-blame: DISAPPOINTMENT. They become REGRET only when the
  reviewer also blames their own past decision ("I should have stuck with Duracell",
  "lesson learned, I should have paid more").
- **"Waste of money"** on its own = DISAPPOINTMENT; with self-blame for the decision
  ("I wasted my money, should have listened to the reviews") = REGRET.
- Self-blame for trusting the brand, for not reading the description or reviews, for ordering the
  wrong item, or for misusing the product ("my fault", "I ordered AAA by mistake") = REGRET.
- Missing a return window: REGRET if the reviewer blames themselves ("I waited too long"),
  DISAPPOINTMENT if they blame the policy.
- Short or uninformative reviews ("ok", "not great"): primary_emotion = NEITHER if no emotion can
  be read; still make a forced_choice.

## Practical
- Use the drop-down menus. Do not change the review_uid column or the item order.
- Expect about 20-30 seconds per review for the required fields, about 1 minute with all fields.
- Save the file under the same name in `annotation/returned/`. It is fine for coder B to code only
  part of the sheet (top to bottom); agreement is computed on the reviews both coders completed.
"""


# --------------------------------------------------------------------------
# targeted strata (defined on the full pool; used for sampling weights)
STRATA = [
    ("S1_C3_regret", "C3 (proposed RAG) says REGRET"),
    ("S2_C4_only_regret", "Speculative RAG says REGRET, C3 says DISAPPOINTMENT"),
    ("S3_R0_only_regret", "Original pipeline says REGRET, C3 and C4 say DISAPPOINTMENT"),
    ("S4_lexicon_only", "Explicit regret phrase, all LLM systems say DISAPPOINTMENT"),
    ("S0_rest", "All other reviews"),
]
TARGET_N = {"S1_C3_regret": 25, "S2_C4_only_regret": 30, "S3_R0_only_regret": 30, "S4_lexicon_only": 15}


def assign_strata(pool: pd.DataFrame) -> pd.Series:
    """Stratum of every pool review, from the stored model outputs (no API calls)."""
    m = pd.read_csv(C.RESULTS_DIR / "03_llm_outputs" / "pool_master.csv")
    lex_f = C.RESULTS_DIR / "05_baselines" / "pool_lexicon_vader.csv"
    lex = pd.read_csv(lex_f)[["review_uid", "lexicon_label"]] if lex_f.exists() else None
    m = m.set_index("review_uid")
    g = lambda c: m[c] if c in m else pd.Series("DISAPPOINTMENT", index=m.index)
    c3, c4, r0 = g("C3_rag_full__label"), g("C4_speculative__label"), g("R0_original_replication__label")
    lx = lex.set_index("review_uid")["lexicon_label"].reindex(m.index) if lex is not None else pd.Series("", index=m.index)
    st = pd.Series("S0_rest", index=m.index)
    st[(lx == "REGRET")] = "S4_lexicon_only"
    st[(r0 == "REGRET")] = "S3_R0_only_regret"
    st[(c4 == "REGRET")] = "S2_C4_only_regret"
    st[(c3 == "REGRET")] = "S1_C3_regret"
    return st.reindex(pool["review_uid"]).fillna("S0_rest").values


def select_annotation_set(pool: pd.DataFrame, targeted: bool = True) -> pd.DataFrame:
    """200 random reviews (the paper's 100 + 100 rating-stratified) + ~100 targeted reviews.

    Every coded review carries a stratum and a weight = N_stratum / n_coded_in_stratum, so
    prevalence and accuracy can be estimated without bias (stratified estimator), while the
    targeted reviews give enough REGRET-like cases to measure regret precision/recall."""
    import data_prep
    s100 = data_prep.original_sample(pool).assign(gold_origin="random_paper_sample_100")
    rest = pool[~pool["review_uid"].isin(s100["review_uid"])]
    frac = C.ANNOTATION_N_NEW / len(rest)
    new = pd.concat([g.sample(n=max(1, round(len(g) * frac)), random_state=C.SEED)
                     for _, g in rest.groupby("reviews.rating")]).assign(gold_origin="random_new_100")
    gold = pd.concat([s100, new])
    pool = pool.assign(stratum=assign_strata(pool)) if targeted else pool.assign(stratum="S0_rest")
    gold = gold.merge(pool[["review_uid", "stratum"]], on="review_uid")
    if targeted:
        left = pool[~pool["review_uid"].isin(gold["review_uid"])]
        picks = []
        for h, n in TARGET_N.items():
            cand = left[left["stratum"] == h]
            if len(cand):
                picks.append(cand.sample(n=min(n, len(cand)), random_state=C.SEED))
        if picks:
            gold = pd.concat([gold, pd.concat(picks).assign(gold_origin="targeted")], ignore_index=True)
    N_h = pool["stratum"].value_counts()
    n_h = gold["stratum"].value_counts()
    gold["stratum_N"] = gold["stratum"].map(N_h)
    gold["stratum_n_coded"] = gold["stratum"].map(n_h)
    gold["weight"] = gold["stratum_N"] / gold["stratum_n_coded"]
    return gold


def _practice_set(pool: pd.DataFrame, exclude: set, n: int = 10) -> pd.DataFrame:
    """10 calibration reviews, not part of the study: 5 random + 5 from the regret-like strata."""
    p = pool[~pool["review_uid"].isin(exclude)].assign(stratum=assign_strata(pool[~pool["review_uid"].isin(exclude)]))
    a = p[p["stratum"] == "S0_rest"].sample(n=n // 2, random_state=7)
    b = p[p["stratum"] != "S0_rest"].sample(n=n - n // 2, random_state=7)
    return pd.concat([a, b]).sample(frac=1, random_state=7)


def _write_workbook(df: pd.DataFrame, path, coder: str, practice: pd.DataFrame | None = None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook(); ws = wb.active; ws.title = "coding"
    cols = ["item", "review_uid", "product", "review_text", "primary_emotion", "forced_choice", "cf_focus",
            "counterfactual", "responsibility", "controllability", "affect", "certainty", "notes"]
    ws.append(cols)
    for i, (_, r) in enumerate(df.iterrows(), start=1):
        ws.append([i, r["review_uid"], str(r.get("name", ""))[:80], str(r["reviews.text"])] + [None] * 9)
    widths = [6, 10, 28, 80, 18, 16, 13, 14, 15, 15, 9, 11, 30]
    for c, w in zip("ABCDEFGHIJKLM", widths):
        ws.column_dimensions[c].width = w
    head = PatternFill("solid", fgColor="DDE6F0"); code = PatternFill("solid", fgColor="FFF7D6")
    req = PatternFill("solid", fgColor="F6C9A8")
    for cell in ws[1]:
        cell.font = Font(bold=True); cell.fill = head
    for col in ("E", "F", "L"):                      # required columns: orange header
        ws[f"{col}1"].fill = req
        ws[f"{col}1"].value = ws[f"{col}1"].value      # names stay unchanged for the parser
    n = len(df) + 1
    for row in ws.iter_rows(min_row=2, max_row=n):
        row[3].alignment = Alignment(wrap_text=True, vertical="top")
        for c in row[4:]:
            c.fill = code
    def dv(options, col):
        v = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        v.error = "Choose from the list"; ws.add_data_validation(v); v.add(f"{col}2:{col}{n}")
    dv(EMOTIONS, "E"); dv(FORCED, "F"); dv(FOCUS, "G")
    for col in "HIJK":
        dv(SCALE, col)
    dv(["LOW", "MEDIUM", "HIGH"], "L")
    ws.freeze_panes = "E2"
    if practice is not None:
        ps = wb.create_sheet("practice")
        ps.append(["item", "review_text", "primary_emotion", "forced_choice", "certainty", "notes"])
        for i, (_, r) in enumerate(practice.iterrows(), start=1):
            ps.append([i, str(r["reviews.text"]), None, None, None, None])
        ps.column_dimensions["B"].width = 90
        for row in ps.iter_rows(min_row=2):
            row[1].alignment = Alignment(wrap_text=True, vertical="top")
        for opts, col in [(EMOTIONS, "C"), (FORCED, "D"), (["LOW", "MEDIUM", "HIGH"], "E")]:
            v = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
            ps.add_data_validation(v); v.add(f"{col}2:{col}{len(practice) + 1}")
    cb = wb.create_sheet("codebook")
    for line in CODEBOOK.splitlines():
        cb.append([line])
    cb.column_dimensions["A"].width = 120
    meta = wb.create_sheet("meta"); meta.append(["coder", coder]); meta.append(["n_items", len(df)])
    wb.save(path)


def build_sheets(pool: pd.DataFrame | None = None, force: bool = False) -> pd.DataFrame:
    """Creates the coder workbooks. Never overwrites existing sheets unless force=True."""
    C.ensure_dirs(); OUT.mkdir(parents=True, exist_ok=True)
    if not force and (C.ANNOT_DIR / "annotation_coder_A.xlsx").exists():
        print("[annotation] coder sheets already exist - not overwritten (use force=True to rebuild)")
        return pd.read_csv(C.DATA_DIR / "annotation_set.csv")
    if pool is None:
        pool = pd.read_csv(C.DATA_DIR / "pool.csv")
    targeted = (C.RESULTS_DIR / "03_llm_outputs" / "pool_master.csv").exists()
    gold = select_annotation_set(pool, targeted=targeted)
    gold[["review_uid", "gold_origin", "stratum", "stratum_N", "stratum_n_coded", "weight",
          "reviews.rating", "primaryCategories"]].to_csv(C.DATA_DIR / "annotation_set.csv", index=False)
    practice = _practice_set(pool, set(gold["review_uid"])) if targeted else None
    for coder, seed in [("A", 1), ("B", 2)]:
        order = gold.sample(frac=1, random_state=seed)   # random order mixes random and targeted items
        _write_workbook(order, C.ANNOT_DIR / f"annotation_coder_{coder}.xlsx", coder, practice)
    pd.DataFrame([{"stratum": k, "description": d,
                   "N_in_pool": int((assign_strata(pool) == k).sum()) if targeted else np.nan,
                   "n_coded": int((gold["stratum"] == k).sum())} for k, d in STRATA]).to_csv(
        OUT / "annotation_strata.csv", index=False)
    (C.ANNOT_DIR / "CODEBOOK.md").write_text(CODEBOOK, encoding="utf-8")
    comp = pd.DataFrame({
        "gold_n": gold["reviews.rating"].value_counts().sort_index(),
        "gold_random_n": gold[gold["gold_origin"] != "targeted"]["reviews.rating"].value_counts().sort_index(),
        "gold_pct": (gold["reviews.rating"].value_counts(normalize=True).sort_index() * 100).round(1),
        "pool_pct": (pool["reviews.rating"].value_counts(normalize=True).sort_index() * 100).round(1)})
    comp.to_csv(OUT / "annotation_set_composition.csv")
    return gold


# --------------------------------------------------------------------------
def _read_coder(coder: str) -> pd.DataFrame | None:
    f = C.ANNOT_RETURNED_DIR / f"annotation_coder_{coder}.xlsx"
    if not f.exists():
        return None
    d = pd.read_excel(f, sheet_name="coding")
    d = d.drop(columns=[c for c in ["item", "product", "review_text"] if c in d])
    for c in ["primary_emotion", "forced_choice", "cf_focus", "certainty"]:
        d[c] = d[c].astype(str).str.strip().str.upper().replace({"NAN": np.nan, "NONE ": "NONE"})
    for c in DIMS:
        d[c] = pd.to_numeric(d[c], errors="coerce")
    return d.set_index("review_uid")


def agreement() -> dict:
    """Inter-rater reliability + adjudication sheet + gold labels."""
    from sklearn.metrics import cohen_kappa_score
    import krippendorff
    import stats_utils as S

    OUT.mkdir(parents=True, exist_ok=True)
    a, b = _read_coder("A"), _read_coder("B")
    if a is None or b is None:
        msg = "Waiting for both coder files in annotation/returned/ (annotation_coder_A.xlsx, annotation_coder_B.xlsx)"
        print(msg); return {"status": msg}
    m = a.join(b, lsuffix="_A", rsuffix="_B", how="outer")
    coded_A = m["forced_choice_A"].notna(); coded_B = m["forced_choice_B"].notna()
    rows = []
    for c in ["primary_emotion", "forced_choice", "cf_focus"]:
        x = m[[f"{c}_A", f"{c}_B"]].dropna()
        k = cohen_kappa_score(x[f"{c}_A"], x[f"{c}_B"])
        lo, hi = S.bootstrap_ci(lambda i: cohen_kappa_score(x[f"{c}_A"].values[i], x[f"{c}_B"].values[i]), len(x))
        rows.append({"field": c, "n": len(x), "pct_agreement": round((x[f"{c}_A"] == x[f"{c}_B"]).mean(), 3),
                     "cohen_kappa": round(k, 3), "kappa_ci95": f"[{lo:.3f}, {hi:.3f}]",
                     "krippendorff_alpha": round(krippendorff.alpha(
                         reliability_data=[pd.factorize(pd.concat([x[f'{c}_A'], x[f'{c}_B']]))[0][:len(x)],
                                           pd.factorize(pd.concat([x[f'{c}_A'], x[f'{c}_B']]))[0][len(x):]],
                         level_of_measurement="nominal"), 3)})
    for c in DIMS:
        x = m[[f"{c}_A", f"{c}_B"]].dropna()
        rows.append({"field": c, "n": len(x), "pct_agreement": round((x[f"{c}_A"] == x[f"{c}_B"]).mean(), 3),
                     "pct_within_1": round(((x[f"{c}_A"] - x[f"{c}_B"]).abs() <= 1).mean(), 3),
                     "weighted_kappa_quadratic": round(cohen_kappa_score(x[f"{c}_A"].astype(int), x[f"{c}_B"].astype(int), weights="quadratic"), 3),
                     "krippendorff_alpha": round(krippendorff.alpha(reliability_data=[x[f"{c}_A"].values, x[f"{c}_B"].values],
                                                                    level_of_measurement="ordinal"), 3),
                     "icc2_1": round(S.icc2_1(x[f"{c}_A"].values, x[f"{c}_B"].values), 3)})
    irr = pd.DataFrame(rows)
    irr.to_csv(OUT / "inter_rater_agreement.csv", index=False)

    # adjudication workbook (disagreements on the categorical fields)
    pool = pd.read_csv(C.DATA_DIR / "pool.csv").set_index("review_uid")
    both = m[coded_A & coded_B]
    dis = both[(both["primary_emotion_A"] != both["primary_emotion_B"]) | (both["forced_choice_A"] != both["forced_choice_B"])]
    adj_f = C.ANNOT_DIR / "adjudication.xlsx"
    adj_ret = C.ANNOT_RETURNED_DIR / "adjudication.xlsx"
    if not adj_ret.exists():
        adj = dis[["primary_emotion_A", "primary_emotion_B", "forced_choice_A", "forced_choice_B"]].copy()
        adj.insert(0, "review_text", pool.loc[adj.index, "reviews.text"])
        adj["primary_emotion_FINAL"] = ""; adj["forced_choice_FINAL"] = ""
        adj.to_excel(adj_f)

    # gold labels: both agree -> that value; disagree -> adjudicated value (else unresolved);
    # coded by only one coder -> that coder's value (flagged single_coded)
    gold = pd.DataFrame(index=m.index)
    for c in ["primary_emotion", "forced_choice", "cf_focus"]:
        A_, B_ = m[f"{c}_A"], m[f"{c}_B"]
        gold[c] = np.where(A_.notna() & B_.notna(), np.where(A_ == B_, A_, np.nan),
                           np.where(A_.notna(), A_, B_))
    gold["double_coded"] = coded_A & coded_B
    gold["single_coded"] = coded_A ^ coded_B
    gold = gold[coded_A | coded_B]
    if adj_ret.exists():
        r = pd.read_excel(adj_ret, index_col=0)
        for c in ["primary_emotion", "forced_choice"]:
            fin = r[f"{c}_FINAL"].astype(str).str.strip().str.upper()
            fin = fin[fin.isin(EMOTIONS + FORCED)]
            gold.loc[fin.index, c] = fin
    for c in DIMS:
        gold[f"human_{c}"] = m.loc[gold.index, [f"{c}_A", f"{c}_B"]].mean(axis=1)
    gold["adjudicated"] = adj_ret.exists()
    aset = pd.read_csv(C.DATA_DIR / "annotation_set.csv").set_index("review_uid")
    for col in ["gold_origin", "stratum", "weight"]:
        if col in aset:
            gold[col] = aset[col].reindex(gold.index)
    gold.reset_index().to_csv(C.DATA_DIR / "gold_labels.csv", index=False)
    summary = {"n_items_coded": int(len(gold)), "n_double_coded": int(gold["double_coded"].sum()),
               "n_disagreements": len(dis),
               "gold_forced_choice_counts": gold["forced_choice"].value_counts(dropna=False).to_dict(),
               "gold_primary_counts": gold["primary_emotion"].value_counts(dropna=False).to_dict()}
    (OUT / "gold_summary.json").write_text(json.dumps(summary, indent=2, default=str), encoding="utf-8")
    return summary


if __name__ == "__main__":
    import sys
    g = build_sheets(force="--force" in sys.argv)
    print(g["gold_origin"].value_counts(), g["reviews.rating"].value_counts())
