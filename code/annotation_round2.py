"""
Round-2 human annotation (second-round review: Reviewers 2 and 3).

Builds the set and the two coder workbooks, and scores returned workbooks.

Set (397 rows per coder, each coder in a different random order):
  * 285 reviews of round 1 (200 random + 85 targeted), unchanged
  * 100 extra targeted reviews, so that enough regret cases exist to estimate
    regret precision/recall (Reviewer 3): every remaining review C3 labels regret,
    13 more "original-pipeline-only" reviews, and the rest "Speculative-RAG-only"
  * 12 attention checks with an unambiguous answer under the codebook (6 regret,
    6 disappointment), mixed in unmarked and excluded from every analysis

Stratum weights are recomputed so prevalence and accuracy stay unbiased.

Outputs
  data/annotation_set_round2.csv                       study set with strata and weights
  data/attention_checks_round2.csv                     answer key (do NOT give to coders)
  annotation/round2/annotation_round2_coder_1.xlsx     workbook, coder 1
  annotation/round2/annotation_round2_coder_2.xlsx     workbook, coder 2
Returned workbooks go to annotation/round2/returned/ under the same names.
"""
from __future__ import annotations

import numpy as np
import pandas as pd

import annotation as A
import config as C

R2_DIR = C.ANNOT_DIR / "round2"
R2_RETURNED = R2_DIR / "returned"
OUT = C.RESULTS_DIR / "04_human_annotation_round2"
SEED = 2026
N_EXTRA = 100
EXTRA = {"S3_R0_only_regret": 13}
ATTENTION = {  # review_uid -> expected forced choice (unambiguous under the codebook)
    "P0361": "REGRET", "P0788": "REGRET", "P1201": "REGRET",
    "P1695": "REGRET", "P1752": "REGRET", "P1848": "REGRET",
    "P0139": "DISAPPOINTMENT", "P0446": "DISAPPOINTMENT", "P0465": "DISAPPOINTMENT",
    "P0466": "DISAPPOINTMENT", "P0740": "DISAPPOINTMENT", "P1112": "DISAPPOINTMENT",
}
ATTENTION_PASS = 10  # of 12


def build(force: bool = False) -> pd.DataFrame:
    R2_DIR.mkdir(parents=True, exist_ok=True); R2_RETURNED.mkdir(parents=True, exist_ok=True)
    w1 = R2_DIR / "annotation_round2_coder_1.xlsx"
    if w1.exists() and not force:
        print("[round2] workbooks exist - not overwritten (force=True to rebuild)")
        return pd.read_csv(C.DATA_DIR / "annotation_set_round2.csv")
    pool = pd.read_csv(C.DATA_DIR / "pool.csv")
    pool = pool.assign(stratum=A.assign_strata(pool))
    r1 = pd.read_csv(C.DATA_DIR / "annotation_set.csv")
    practice = A._practice_set(pool.drop(columns="stratum"), set(r1["review_uid"]))
    taken = set(r1["review_uid"]) | set(practice["review_uid"]) | set(ATTENTION)
    left = pool[~pool["review_uid"].isin(taken)]
    s1 = left[left["stratum"] == "S1_C3_regret"]          # take every remaining C3-regret review
    n_s2 = N_EXTRA - len(s1) - EXTRA["S3_R0_only_regret"]  # fill the rest from Speculative-RAG-only
    extra = pd.concat([s1,
                       left[left["stratum"] == "S2_C4_only_regret"].sample(n=n_s2, random_state=SEED),
                       left[left["stratum"] == "S3_R0_only_regret"].sample(n=EXTRA["S3_R0_only_regret"], random_state=SEED)])
    extra = extra.assign(gold_origin="targeted_round2")
    base = pool.merge(r1[["review_uid", "gold_origin"]], on="review_uid")
    study = pd.concat([base, extra], ignore_index=True)
    N_h = pool["stratum"].value_counts(); n_h = study["stratum"].value_counts()
    study["stratum_N"] = study["stratum"].map(N_h); study["stratum_n_coded"] = study["stratum"].map(n_h)
    study["weight"] = study["stratum_N"] / study["stratum_n_coded"]
    study[["review_uid", "gold_origin", "stratum", "stratum_N", "stratum_n_coded", "weight",
           "reviews.rating", "primaryCategories"]].to_csv(C.DATA_DIR / "annotation_set_round2.csv", index=False)
    att = pool[pool["review_uid"].isin(ATTENTION)].assign(expected=lambda d: d["review_uid"].map(ATTENTION))
    att[["review_uid", "expected", "reviews.text"]].to_csv(C.DATA_DIR / "attention_checks_round2.csv", index=False)
    rows = pd.concat([study, att.drop(columns="expected")], ignore_index=True)
    for coder, seed in [("1", 11), ("2", 22)]:
        A._write_workbook(rows.sample(frac=1, random_state=seed), R2_DIR / f"annotation_round2_coder_{coder}.xlsx",
                          f"round2-{coder}", practice)
        _add_declaration(R2_DIR / f"annotation_round2_coder_{coder}.xlsx")
    return study


def _add_declaration(path):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font
    wb = load_workbook(path)
    ws = wb.create_sheet("declaration", 0)
    lines = [
        ("Annotator declaration (complete before you start, sign when you finish)", True),
        ("", False),
        ("1. I coded every review myself, by reading it, without help from any AI tool, chatbot, script or other person.", False),
        ("2. I did not discuss specific reviews with the other annotator before we both submitted.", False),
        ("3. I did not look at star ratings, model outputs or any analysis of these reviews.", False),
        ("4. I followed the coding manual (Annotation_Instructions document and the codebook sheet).", False),
        ("", False),
        ("Name:", True), ("Date started:", True), ("Date finished:", True),
        ("Approximate total hours spent:", True), ("Signature (type your full name):", True),
    ]
    for text, bold in lines:
        ws.append([text]); ws.cell(ws.max_row, 1).font = Font(bold=bold)
    ws.column_dimensions["A"].width = 110; ws.column_dimensions["B"].width = 40
    for r in range(1, ws.max_row + 1):
        ws.cell(r, 1).alignment = Alignment(wrap_text=True)
    wb.save(path)


# --------------------------------------------------------------------------
def _read(coder: str) -> pd.DataFrame | None:
    f = R2_RETURNED / f"annotation_round2_coder_{coder}.xlsx"
    if f.exists():
        return pd.read_excel(f, sheet_name="coding").drop(columns=["item", "product", "review_text"], errors="ignore")
    g = R2_RETURNED / f"annotation_round2_coder_{coder}.csv"     # de-identified copy (public repository)
    if g.exists():
        return pd.read_csv(g).drop(columns=["item", "product", "review_text"], errors="ignore")
    return None


def provenance(coder: str) -> dict:
    """Checks that flagged the round-1 files: writing application, notes on every row."""
    import zipfile, re
    f = R2_RETURNED / f"annotation_round2_coder_{coder}.xlsx"
    if not f.exists():
        return {"coder": coder, "application": "not available (de-identified CSV)"}
    z = zipfile.ZipFile(f)
    names = z.namelist()
    app_xml = z.read("docProps/app.xml").decode(errors="ignore") if "docProps/app.xml" in names else ""
    app = re.search(r"<Application>(.*?)</Application>", app_xml)
    if not app and "GoogleSheetsCustomData" in z.read("xl/workbook.xml").decode(errors="ignore"):
        app = re.search("(.+)", "Google Sheets export")
    core = z.read("docProps/core.xml").decode(errors="ignore")
    mod = re.search(r"<cp:lastModifiedBy>(.*?)</cp:lastModifiedBy>", core)
    d = _read(coder)
    return {"coder": coder, "application": app.group(1) if app else None,
            "last_modified_by": mod.group(1) if mod else None,
            "pct_rows_with_notes": round(100 * d["notes"].notna().mean(), 1),
            "n_rows_coded": int(d["forced_choice"].notna().sum())}


def score() -> dict:
    """Attention checks, provenance, inter-coder agreement, gold labels. Run after both files return."""
    from sklearn.metrics import cohen_kappa_score
    from automated_annotation import krippendorff_nominal
    OUT.mkdir(parents=True, exist_ok=True)
    a, b = _read("1"), _read("2")
    if a is None or b is None:
        print("[round2] waiting for both returned workbooks in", R2_RETURNED); return {}
    att = pd.read_csv(C.DATA_DIR / "attention_checks_round2.csv")
    res = {"provenance": [provenance("1"), provenance("2")], "attention": {}}
    for k, d in [("1", a), ("2", b)]:
        j = att.merge(d, on="review_uid")
        res["attention"][k] = int((j["forced_choice"] == j["expected"]).sum())
    m = a.merge(b, on="review_uid", suffixes=("_1", "_2"))
    m = m[~m["review_uid"].isin(att["review_uid"])]
    both = m.dropna(subset=["forced_choice_1", "forced_choice_2"])
    res["n_both"] = len(both)
    for f in ["forced_choice", "primary_emotion", "cf_focus"]:
        d = both.dropna(subset=[f"{f}_1", f"{f}_2"])
        if d.empty:
            continue
        x, y = d[f"{f}_1"].astype(str).tolist(), d[f"{f}_2"].astype(str).tolist()
        res[f"kappa_{f}"] = round(cohen_kappa_score(x, y), 3)
        res[f"alpha_{f}"] = round(krippendorff_nominal(x, y), 3)
        res[f"pct_agree_{f}"] = round(100 * float(np.mean(np.array(x) == np.array(y))), 1)
    dis = both[both["forced_choice_1"] != both["forced_choice_2"]]
    dis.to_csv(OUT / "disagreements_for_adjudication.csv", index=False)
    pd.Series(res).to_json(OUT / "round2_scoring.json", indent=1)
    print(pd.Series(res).to_string())
    return res


if __name__ == "__main__":
    s = build()
    print(s["stratum"].value_counts().to_string(), len(s))
